import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_elegant_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Warm, Simple & Elegant Color Palette (Matching User's Reference Screenshot)
    COLOR_BG = RGBColor(250, 247, 242)        # #FAF7F2 Warm Cream / Ivory
    COLOR_TITLE = RGBColor(56, 42, 36)        # #382A24 Deep Espresso Brown
    COLOR_BODY = RGBColor(89, 73, 65)         # #594941 Warm Charcoal Brown
    COLOR_MUTED = RGBColor(145, 131, 122)     # #91837A Soft Sandstone
    
    # Warm Earthy Accents
    COLOR_TERRACOTTA = RGBColor(217, 123, 78) # #D97B4E Burnt Orange / Terracotta Accent
    COLOR_SAND_CARD = RGBColor(243, 239, 233) # #F3EFE9 Warm Beige Card
    COLOR_CARD_BORDER = RGBColor(228, 220, 210) # #E4DCD2 Subtle Card Border
    COLOR_BLUSH_BG = RGBColor(245, 238, 232)  # #F5EEE8 Very Soft Warm Blush
    COLOR_GREEN_MINT = RGBColor(76, 138, 107) # #4C8A6B Muted Olive Green for Real Data
    COLOR_RED_CLAY = RGBColor(194, 76, 60)    # #C24C3C Muted Clay Red for Previous Test Data

    FONT_FAMILY = "Helvetica Neue"

    def set_slide_bg(slide, color=COLOR_BG):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_organic_accents(slide):
        # Soft subtle organic oval shape in top-left or bottom-right
        oval1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-0.8), Inches(-0.8), Inches(3.2), Inches(3.2))
        oval1.fill.solid()
        oval1.fill.fore_color.rgb = COLOR_BLUSH_BG
        oval1.line.fill.background()

        # Elegant thin terracotta curved line representation at top right
        line1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(0.4), Inches(2.2), Inches(0.04))
        line1.fill.solid()
        line1.fill.fore_color.rgb = COLOR_TERRACOTTA
        line1.line.fill.background()

        # Soft oval in bottom right
        oval2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.2), Inches(5.6), Inches(3.0), Inches(2.8))
        oval2.fill.solid()
        oval2.fill.fore_color.rgb = COLOR_SAND_CARD
        oval2.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=COLOR_SAND_CARD, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    def add_slide_header(slide, category_text, title_text, subtitle_text):
        # Subtle Category Kicker
        k_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.3))
        ktf = k_box.text_frame
        ktf.word_wrap = True
        ktf.margin_left = ktf.margin_top = ktf.margin_right = ktf.margin_bottom = 0
        kp = ktf.paragraphs[0]
        kp.text = category_text.upper()
        kp.font.name = FONT_FAMILY
        kp.font.size = Pt(9.5)
        kp.font.bold = True
        kp.font.color.rgb = COLOR_TERRACOTTA

        # Main Title
        t_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.85), Inches(11.5), Inches(0.6))
        ttf = t_box.text_frame
        ttf.word_wrap = True
        ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0
        tp = ttf.paragraphs[0]
        tp.text = title_text
        tp.font.name = FONT_FAMILY
        tp.font.size = Pt(22)
        tp.font.bold = True
        tp.font.color.rgb = COLOR_TITLE

        # Subtitle in plain, human English
        if subtitle_text:
            s_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.48), Inches(11.5), Inches(0.35))
            stf = s_box.text_frame
            stf.word_wrap = True
            stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
            sp = stf.paragraphs[0]
            sp.text = subtitle_text
            sp.font.name = FONT_FAMILY
            sp.font.size = Pt(12)
            sp.font.color.rgb = COLOR_BODY

        # Subtle thin divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.9), Inches(11.533), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_CARD_BORDER
        line.line.fill.background()

    # =========================================================================
    # SLIDE 1: Title Slide (Simple & Elegant, exactly matching reference image)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)
    add_organic_accents(s1)

    # Big Central Elegant Title Box
    title_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(3.2))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "FINAL YEAR PROJECT DEFENSE"
    p0.font.name = FONT_FAMILY
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TERRACOTTA
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf1.add_paragraph()
    p1.text = "KUBEMIND"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TITLE
    p1.alignment = PP_ALIGN.CENTER
    p1.space_before = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "AN AI COPILOT FOR KUBERNETES INCIDENT DIAGNOSIS"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_BODY
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)

    # Decorative warm line under title
    c_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.666), Inches(4.3), Inches(2.0), Inches(0.03))
    c_line.fill.solid()
    c_line.fill.fore_color.rgb = COLOR_TERRACOTTA
    c_line.line.fill.background()

    # Presenter Card at the bottom
    p_card = add_card(s1, Inches(3.666), Inches(4.8), Inches(6.0), Inches(1.3), bg_color=COLOR_SAND_CARD, border_color=COLOR_CARD_BORDER)
    tb_pres = s1.shapes.add_textbox(Inches(3.8), Inches(4.95), Inches(5.7), Inches(1.0))
    ptf = tb_pres.text_frame
    ptf.word_wrap = True

    pp1 = ptf.paragraphs[0]
    pp1.text = "Student: Harmit Jetani  •  B.Tech Computer Science & Engineering"
    pp1.font.name = FONT_FAMILY
    pp1.font.size = Pt(11)
    pp1.font.bold = True
    pp1.font.color.rgb = COLOR_TITLE
    pp1.alignment = PP_ALIGN.CENTER

    pp2 = ptf.add_paragraph()
    pp2.text = "Focus: Feasibility Evaluation & Real Kubernetes Cluster Telemetry"
    pp2.font.name = FONT_FAMILY
    pp2.font.size = Pt(10)
    pp2.font.color.rgb = COLOR_MUTED
    pp2.alignment = PP_ALIGN.CENTER
    pp2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: Problem Statement (Explained in Simple, Everyday Words)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_organic_accents(s2)
    add_slide_header(s2, "01 / Why We Built This", "The Problem: Why Kubernetes Is Hard to Troubleshoot", "When an application fails inside Kubernetes, finding out why takes too long and requires manual work.")

    cards2 = [
        ("Too Many Confusing Alerts", "When one background service crashes, 20 other services fail at the same time. The person on call gets flooded with error messages and cannot quickly tell which pod caused the issue."),
        ("Too Many Terminal Commands", "To find the root cause, an engineer must manually open terminal windows, run 'kubectl logs', 'kubectl describe', check Prometheus graphs, and decode exit codes. This often takes 30 to 45 minutes."),
        ("Dashboards Only Show Graphs", "Tools like Grafana show nice graphs that an error happened, but they cannot tell you why it happened or what command to run to fix it. A human still has to figure it out.")
    ]

    for i, (head, desc) in enumerate(cards2):
        cx = Inches(0.9 + i * 3.95)
        add_card(s2, cx, Inches(2.2), Inches(3.7), Inches(4.6))
        
        tb = s2.shapes.add_textbox(cx + Inches(0.3), Inches(2.45), Inches(3.1), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = f"0{i+1}"
        p_num.font.name = FONT_FAMILY
        p_num.font.size = Pt(22)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_TERRACOTTA

        p_h = tf.add_paragraph()
        p_h.text = head
        p_h.font.name = FONT_FAMILY
        p_h.font.size = Pt(14)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_TITLE
        p_h.space_before = Pt(8)

        p_b = tf.add_paragraph()
        p_b.text = desc
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = COLOR_BODY
        p_b.space_before = Pt(10)

    # =========================================================================
    # SLIDE 3: Directly Addressing the External Mentor's Feedback
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_organic_accents(s3)
    add_slide_header(s3, "02 / Response to Reviewer Feedback", "Addressing the Mentor's Remarks Directly", "How we improved the project based on the review: proving feasibility and using 100% real data.")

    # Mentor Remark Callout Box
    add_card(s3, Inches(0.9), Inches(2.2), Inches(11.533), Inches(1.2), bg_color=COLOR_BLUSH_BG, border_color=COLOR_TERRACOTTA)
    tb_q = s3.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(11.0), Inches(1.0))
    qtf = tb_q.text_frame
    qtf.word_wrap = True
    qp1 = qtf.paragraphs[0]
    qp1.text = "REVIEWER FEEDBACK (PREVIOUS SCORE: 37 / 50)"
    qp1.font.name = FONT_FAMILY
    qp1.font.size = Pt(9.5)
    qp1.font.bold = True
    qp1.font.color.rgb = COLOR_TERRACOTTA

    qp2 = qtf.add_paragraph()
    qp2.text = '"The idea is good, but still need to work a lot on the system’s feasibility and focus more on using real data rather than test data for testing."'
    qp2.font.name = FONT_FAMILY
    qp2.font.size = Pt(12.5)
    qp2.font.bold = True
    qp2.font.color.rgb = COLOR_TITLE
    qp2.space_before = Pt(3)

    # Two Simple Response Columns
    resp_columns = [
        ("Action 1: Proving Feasibility (Can it run?)", [
            ("Runs on Normal Laptop", "The entire project (Frontend, Backend, Database, AI Service) runs using less than 500MB RAM. You don't need expensive cloud servers."),
            ("No Heavy Cluster Agents", "We do not install heavy programs inside Kubernetes nodes. We talk directly to the Kubernetes API, so there is zero slowdown on your apps."),
            ("Fast Responses", "Dashboard and live telemetry load in less than 1 second. The AI analyzes crash logs and answers in under 4 seconds."),
            ("Safe for Humans", "The AI gives suggestions and shows the exact command to run. It does not delete or change your cluster without your approval.")
        ]),
        ("Action 2: Using Real Data (Where data comes from)", [
            ("Deleted All Fake Data", "We completely removed static JSON files and fake mock generators. 100% of the data on screen comes from a live system."),
            ("Live Kubernetes Cluster", "Connected directly to a real Kubernetes cluster running on Docker Desktop (v1.32) with live pods and nodes."),
            ("Real Prometheus Telemetry", "Prometheus actively scrapes real CPU and memory usage from live containers every few seconds."),
            ("Real Crash Testing", "We tested using real container crashes: memory stress reaching real Linux Out-Of-Memory limits (Exit Code 137).")
        ])
    ]

    for i, (title, items) in enumerate(resp_columns):
        cx = Inches(0.9 + i * 5.9)
        add_card(s3, cx, Inches(3.6), Inches(5.633), Inches(3.4))

        tb = s3.shapes.add_textbox(cx + Inches(0.3), Inches(3.75), Inches(5.0), Inches(3.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE

        for b_title, b_desc in items:
            bp = tf.add_paragraph()
            bp.text = f"•  {b_title}: {b_desc}"
            bp.font.name = FONT_FAMILY
            bp.font.size = Pt(10)
            bp.font.color.rgb = COLOR_BODY
            bp.space_before = Pt(6)

    # =========================================================================
    # SLIDE 4: Where Does the Data Come From? (Core Reviewer Question)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_organic_accents(s4)
    add_slide_header(s4, "03 / Data Sources", "Where Does the Project Data Come From?", "A clear look at the 4 live data sources used by KubeMind—no fake data.")

    data_sources = [
        ("1. Pod & Node State", "Kubernetes API Server", "The backend talks directly to the official Kubernetes API on port 6443. It fetches the live list of running pods, node health, restart counts, and namespaces in real time."),
        ("2. CPU & Memory Usage", "Prometheus Scraper", "A real Prometheus container (port 9090) measures the actual CPU and memory used by the containers. The graphs on the dashboard update from live PromQL queries."),
        ("3. Container Crash Logs", "Linux / Docker Runtime", "When an application crashes, KubeMind reads the real terminal output (stdout and stderr) directly from the container via Kubernetes. It sees the exact fatal error message."),
        ("4. System Warnings & Events", "Kubernetes Event Feed", "Kubernetes creates event messages whenever a pod restarts or runs out of memory. KubeMind captures these Warning events live and turns them into incident cards.")
    ]

    for i, (title, source, desc) in enumerate(data_sources):
        bx = Inches(0.9 + (i % 2) * 5.9)
        by = Inches(2.2 + (i // 2) * 2.4)
        add_card(s4, bx, by, Inches(5.633), Inches(2.15))

        tb = s4.shapes.add_textbox(bx + Inches(0.3), by + Inches(0.2), Inches(5.0), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE

        p_src = tf.add_paragraph()
        p_src.text = f"Source: {source}"
        p_src.font.name = FONT_FAMILY
        p_src.font.size = Pt(10)
        p_src.font.bold = True
        p_src.font.color.rgb = COLOR_TERRACOTTA
        p_src.space_before = Pt(2)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = COLOR_BODY
        p_desc.space_before = Pt(6)

    # =========================================================================
    # SLIDE 5: Side-by-Side: Real Data vs. Fake Test Data
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_organic_accents(s5)
    add_slide_header(s5, "04 / Verification Matrix", "Comparison: What Changed from Test Data to Real Data", "An honest, side-by-side comparison showing how every piece of data is now real.")

    add_card(s5, Inches(0.9), Inches(2.2), Inches(11.533), Inches(4.7))
    tb_t = s5.shapes.add_textbox(Inches(1.15), Inches(2.35), Inches(11.0), Inches(4.3))
    ttf = tb_t.text_frame
    ttf.word_wrap = True

    # Simple Header
    hp = ttf.paragraphs[0]
    hp.text = f"{'PROJECT FEATURE':<24}  |  {'BEFORE (PROTOTYPE TEST DATA)':<38}  |  {'NOW (100% REAL CLUSTER DATA)'}"
    hp.font.name = FONT_FAMILY
    hp.font.size = Pt(10)
    hp.font.bold = True
    hp.font.color.rgb = COLOR_TERRACOTTA

    matrix_rows = [
        ("Pod & Workload List", "Hardcoded JSON with fake names like 'auth-service-7f89b'", "Live pods running on Docker Desktop: 'memory-crash-pod', 'live-demo-app', and 'coredns'"),
        ("CPU & Memory Metrics", "Random math numbers generated on a timer", "Actual CPU & RAM measured by Prometheus from real Linux container cgroups"),
        ("Container Crash Logs", "Pre-written text strings saved in a database table", "Real log output streamed directly from container terminal ('stress: dispatching hogs...', 'exit 1')"),
        ("Incident Alerts", "Dummy rows inserted during database setup", "Triggered automatically only when Kubernetes marks a pod as CrashLoopBackOff or OOMKilled"),
        ("Fixing the Issue", "Clicking a button showed a fake success popup", "AI generates a real kubectl command that you run in your terminal to fix the live pod")
    ]

    for feature, before, now in matrix_rows:
        p_f = ttf.add_paragraph()
        p_f.text = feature
        p_f.font.name = FONT_FAMILY
        p_f.font.size = Pt(11)
        p_f.font.bold = True
        p_f.font.color.rgb = COLOR_TITLE
        p_f.space_before = Pt(8)

        p_b = ttf.add_paragraph()
        p_b.text = f"   Before: {before}"
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = COLOR_RED_CLAY

        p_n = ttf.add_paragraph()
        p_n.text = f"   Now:    {now}"
        p_n.font.name = FONT_FAMILY
        p_n.font.size = Pt(10)
        p_n.font.bold = True
        p_n.font.color.rgb = COLOR_GREEN_MINT

    # =========================================================================
    # SLIDE 6: How the AI Diagnoses Problems (Simple 4-Step Explanation)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_organic_accents(s6)
    add_slide_header(s6, "05 / How the AI Works", "How KubeMind Diagnoses a Crash in 4 Simple Steps", "Instead of guessing, the AI follows four clear steps to read the real logs and provide a fix.")

    steps6 = [
        ("Step 1: Notice the Problem", "The operator types a message in chat, or selects a failing pod from the dashboard. The AI identifies which pod and namespace to investigate."),
        ("Step 2: Read Live Logs", "The AI uses a safe tool to read the last 50 lines of logs from the crashing container. It also checks recent Kubernetes Warning events."),
        ("Step 3: Find the Cause", "The AI parses the logs. It checks Linux exit codes: Exit 137 means Out-Of-Memory, while Exit 1 means an application crash or database failure."),
        ("Step 4: Suggest the Fix", "The AI explains the cause in plain English and gives a safe, ready-to-copy command (like 'kubectl set resources ...' to give more memory).")
    ]

    for i, (title, desc) in enumerate(steps6):
        cx = Inches(0.9 + i * 2.95)
        add_card(s6, cx, Inches(2.2), Inches(2.75), Inches(4.6))

        tb = s6.shapes.add_textbox(cx + Inches(0.25), Inches(2.45), Inches(2.25), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = f"0{i+1}"
        p_num.font.name = FONT_FAMILY
        p_num.font.size = Pt(20)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_TERRACOTTA

        p_h = tf.add_paragraph()
        p_h.text = title
        p_h.font.name = FONT_FAMILY
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_TITLE
        p_h.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = COLOR_BODY
        p_d.space_before = Pt(10)

    # =========================================================================
    # SLIDE 7: How We Tested with Real Data (Crash Demonstration)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_organic_accents(s7)
    add_slide_header(s7, "06 / Real Demonstration", "How We Tested: Real Crash Scenarios in Kubernetes", "We tested by creating actual pods that intentionally fail, proving the system works on live data.")

    scenarios = [
        ("Test 1: Real Out-Of-Memory Crash (OOMKilled)",
         "We started a container with a 50MB memory limit, and instructed it to allocate 150MB of RAM.",
         [
             ("What happened", "The container tried to allocate 150MB of memory on the machine."),
             ("Linux response", "The Linux kernel noticed the container exceeded its limit and killed it with Exit Code 137."),
             ("KubeMind response", "Within 1 second, KubeMind showed a red notification badge: 'OOMKilled in memory-crash-pod'."),
             ("AI diagnosis", "The AI read the memory limit breach and suggested giving the pod 1024Mi of memory.")
         ]),
        ("Test 2: Real Application Crash (CrashLoopBackOff)",
         "We started a microservice that simulated a database timeout and exited with Exit Code 1.",
         [
             ("What happened", "The app printed 'FATAL ERROR: Connection refused' and terminated."),
             ("Kubernetes response", "Kubernetes tried restarting it, placing it in 'CrashLoopBackOff' status with restart count increasing."),
             ("KubeMind response", "KubeMind highlighted the pod in red on the Pods page and streamed the fatal error in the Log Explorer."),
             ("AI diagnosis", "The AI spotted the database connection failure and suggested checking database credentials.")
         ])
    ]

    for i, (title, intro, bullet_items) in enumerate(scenarios):
        cx = Inches(0.9 + i * 5.9)
        add_card(s7, cx, Inches(2.2), Inches(5.633), Inches(4.7))

        tb = s7.shapes.add_textbox(cx + Inches(0.3), Inches(2.4), Inches(5.0), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE

        p_intro = tf.add_paragraph()
        p_intro.text = intro
        p_intro.font.name = FONT_FAMILY
        p_intro.font.size = Pt(10)
        p_intro.font.color.rgb = COLOR_MUTED
        p_intro.space_before = Pt(4)

        for b_head, b_text in bullet_items:
            bp = tf.add_paragraph()
            bp.text = f"•  {b_head}: {b_text}"
            bp.font.name = FONT_FAMILY
            bp.font.size = Pt(10)
            bp.font.color.rgb = COLOR_BODY
            bp.space_before = Pt(8)

    # =========================================================================
    # SLIDE 8: System Feasibility (Can this actually run in real life?)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_organic_accents(s8)
    add_slide_header(s8, "07 / Feasibility", "System Feasibility: Why This Works in the Real World", "Addressing the mentor's concern about feasibility—proving this is realistic, lightweight, and safe.")

    feas_points = [
        ("1. Runs on Any Normal Laptop", "The entire system runs locally inside Docker. It uses less than 500MB of RAM total across backend, frontend, Prometheus, and AI services. You do not need expensive cloud infrastructure to run it."),
        ("2. Zero Cluster Overhead", "Many monitoring tools require installing heavy software on every single cluster node. KubeMind does not do that. It connects only to the main Kubernetes API, so your workloads run at full speed."),
        ("3. Human Stays in Control (Safety)", "AI can make mistakes if allowed to delete or change things automatically. In KubeMind, the AI only recommends actions. The human engineer reviews the suggestion before running it."),
        ("4. Keeps Running if Tools Go Down", "If Prometheus is temporarily restarting, KubeMind continues running smoothly without crashing the dashboard. The system is resilient to network pauses.")
    ]

    for i, (title, desc) in enumerate(feas_points):
        bx = Inches(0.9 + (i % 2) * 5.9)
        by = Inches(2.2 + (i // 2) * 2.4)
        add_card(s8, bx, by, Inches(5.633), Inches(2.15))

        tb = s8.shapes.add_textbox(bx + Inches(0.3), by + Inches(0.2), Inches(5.0), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = COLOR_BODY
        p_desc.space_before = Pt(6)

    # =========================================================================
    # SLIDE 9: Project Features Available for Review
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_organic_accents(s9)
    add_slide_header(s9, "08 / Live Features", "What Reviewers Can Test and See Right Now", "Everything built in the project is fully working and ready to demonstrate live.")

    features = [
        ("Dashboard Command Center", "Shows real-time CPU, Memory, total running pods, active nodes, and connected cluster status."),
        ("Live Notification Bell", "Displays a red badge count in the top bar as soon as any container crashes or restarts."),
        ("Incident Center & RCA", "Lists all active incidents with failure reason, severity, container crash logs, and investigation status."),
        ("Log Explorer", "Streams actual terminal logs directly from any pod running in the Kubernetes cluster."),
        ("Interactive AI Copilot", "Chat window where operators ask diagnostic questions and receive root-cause explanations with fixes."),
        ("Telemetry Resource Charts", "Interactive graphs showing CPU, Memory, Network, and Disk usage across 1h, 6h, and 24h ranges.")
    ]

    for i, (f_name, f_desc) in enumerate(features):
        col_idx = i % 3
        row_idx = i // 3
        fx = Inches(0.9 + col_idx * 3.95)
        fy = Inches(2.2 + row_idx * 2.4)
        add_card(s9, fx, fy, Inches(3.7), Inches(2.15))

        tb = s9.shapes.add_textbox(fx + Inches(0.25), fy + Inches(0.2), Inches(3.2), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f_name
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE

        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_BODY
        p2.space_before = Pt(6)

    # =========================================================================
    # SLIDE 10: Conclusion (Matching Reference Screenshot Style)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_organic_accents(s10)

    # Simple Header
    add_slide_header(s10, "09 / Summary", "Conclusion & Project Summary", "KubeMind proves that AI-assisted Kubernetes troubleshooting is feasible, fast, and works with real data.")

    # Two Simple Summary Cards
    c1 = add_card(s10, Inches(0.9), Inches(2.2), Inches(5.633), Inches(4.7))
    tb_c1 = s10.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(5.0), Inches(4.3))
    tfc1 = tb_c1.text_frame
    tfc1.word_wrap = True

    p = tfc1.paragraphs[0]
    p.text = "Summary of What We Achieved"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE

    achieve_bullets = [
        ("Answered Mentor Feedback", "No more fake test data. Everything on screen comes from a real Kubernetes cluster and Prometheus scraper."),
        ("Proved Feasibility", "The entire stack is lightweight (< 500MB RAM), fast (sub-second telemetry), and runs easily on a student laptop."),
        ("Full End-to-End System", "Includes Frontend UI, Backend API, Prometheus monitoring, incident tracking, and an AI diagnostic agent."),
        ("Ready for Live Demo", "Reviewers can run 'kubectl run' in terminal, crash a pod, and see the notification appear live on screen.")
    ]
    for h, b in achieve_bullets:
        bp = tfc1.add_paragraph()
        bp.text = f"✔  {h}: {b}"
        bp.font.name = FONT_FAMILY
        bp.font.size = Pt(10.5)
        bp.font.color.rgb = COLOR_BODY
        bp.space_before = Pt(8)

    # Right Card: Thank you & Q&A
    c2 = add_card(s10, Inches(6.8), Inches(2.2), Inches(5.633), Inches(4.7), bg_color=COLOR_BLUSH_BG, border_color=COLOR_TERRACOTTA)
    tb_c2 = s10.shapes.add_textbox(Inches(7.1), Inches(2.8), Inches(5.0), Inches(3.5))
    tfc2 = tb_c2.text_frame
    tfc2.word_wrap = True

    tp = tfc2.paragraphs[0]
    tp.text = "THANK YOU"
    tp.font.name = FONT_FAMILY
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = COLOR_TITLE
    tp.alignment = PP_ALIGN.CENTER

    tp2 = tfc2.add_paragraph()
    tp2.text = "Open for Questions & Live Demonstration"
    tp2.font.name = FONT_FAMILY
    tp2.font.size = Pt(14)
    tp2.font.color.rgb = COLOR_TERRACOTTA
    tp2.alignment = PP_ALIGN.CENTER
    tp2.space_before = Pt(6)

    tp3 = tfc2.add_paragraph()
    tp3.text = "We are ready to demonstrate the live cluster,\ncrash a pod in terminal, and show the AI diagnosis on screen."
    tp3.font.name = FONT_FAMILY
    tp3.font.size = Pt(11)
    tp3.font.color.rgb = COLOR_BODY
    tp3.alignment = PP_ALIGN.CENTER
    tp3.space_before = Pt(16)

    output_path = "/Users/harmitjetani/Documents/GitHub/kubmind/KubeMind_Project_Presentation.pptx"
    prs.save(output_path)
    print(f"Simple, elegant presentation saved to: {output_path}")

if __name__ == "__main__":
    build_elegant_presentation()
